"""Integration: образцы шаблонов из templates_samples/ (EPIC-011 Step 6)."""

from __future__ import annotations

import hashlib
from pathlib import Path

import pytest
from openpyxl import load_workbook

from domain.template_markers import marker_catalog_markdown
from reports.excel_template import archive_upload, generate_excel_report, validate_archived
from reports.pdf_export import pdf_contains_text
from reports.pdf_template import archive_pdf_upload, generate_pdf_report, validate_pdf

_SAMPLES = Path(__file__).resolve().parents[2] / "templates_samples"
_EXCEL = _SAMPLES / "sample_report.xlsx"
_BRANCH_SUMMARY = _SAMPLES / "branch_summary_report.xlsx"
_PDF = _SAMPLES / "sample_report.pdf"
_MANIFEST = _SAMPLES / "sample_report.regions.json"
_GUIDE = Path(__file__).resolve().parents[2] / "docs" / "report-templates-guide.md"


@pytest.mark.acceptance
def test_sample_excel_validates() -> None:
    validate_archived(_EXCEL)


@pytest.mark.acceptance
def test_sample_excel_validate_generate_pipeline(tmp_path: Path) -> None:
    archive = tmp_path / "archive.xlsx"
    archived = archive_upload(_EXCEL, archive)
    before = hashlib.sha256(_EXCEL.read_bytes()).hexdigest()
    out = tmp_path / "out.xlsx"
    generate_excel_report(
        archived,
        out,
        scalars={
            "report.title": "Справка о доступности",
            "report.period_from": "2026-08-01",
            "report.period_to": "2026-08-31",
        },
        row_records=[
            {
                "employee.full_name": "Сидоров Алексей",
                "employee.position": "Инженер",
                "employee.branch": "Центральный",
            },
            {
                "employee.full_name": "Козлова Мария",
                "employee.position": "Аналитик",
                "employee.branch": "Северный",
            },
        ],
    )
    assert hashlib.sha256(_EXCEL.read_bytes()).hexdigest() == before
    book = load_workbook(out)
    sheet = book.active
    assert sheet is not None
    assert sheet["A2"].value == "Справка о доступности"
    assert sheet["B6"].value == "Сидоров Алексей"
    assert sheet["B7"].value == "Козлова Мария"
    book.close()


@pytest.mark.acceptance
def test_sample_pdf_validates() -> None:
    validate_pdf(_PDF, manifest_path=_MANIFEST)


@pytest.mark.acceptance
def test_sample_pdf_validate_generate_pipeline(tmp_path: Path) -> None:
    archive = tmp_path / "archive.pdf"
    archived = archive_pdf_upload(_PDF, archive, manifest_source=_MANIFEST)
    before = hashlib.sha256(_PDF.read_bytes()).hexdigest()
    out = tmp_path / "out.pdf"
    generate_pdf_report(
        archived,
        out,
        {
            "report.date": "30.08.2026",
            "employee.full_name": "Сидоров Алексей",
        },
    )
    assert hashlib.sha256(_PDF.read_bytes()).hexdigest() == before
    assert pdf_contains_text(out, "Сидоров")


@pytest.mark.acceptance
def test_branch_summary_sample_validates() -> None:
    validate_archived(_BRANCH_SUMMARY)


@pytest.mark.acceptance
def test_branch_summary_sample_generate_pipeline(tmp_path: Path) -> None:
    archive = tmp_path / "archive.xlsx"
    archived = archive_upload(_BRANCH_SUMMARY, archive)
    out = tmp_path / "out.xlsx"
    generate_excel_report(
        archived,
        out,
        scalars={
            "report.title": "Сводка по филиалу",
            "report.date": "19.09.2026",
            "employee.branch": "Центральный",
            "report.branch_total": "37",
            "report.branch_absent": "7",
        },
        row_records=[
            {
                "employee.division": "Продажи",
                "report.department_total": "25",
                "report.department_absent": "5",
                "report.vacation_employees": "Иванов И.И.; Петров П.П.",
                "report.sick_leave_employees": "Сидоров С.С.",
                "report.business_trip_employees": "Смирнов А.А.; Кузнецов К.К.",
            },
            {
                "employee.division": "Бухгалтерия",
                "report.department_total": "12",
                "report.department_absent": "2",
                "report.vacation_employees": "Орлова О.О.",
                "report.sick_leave_employees": "Волкова В.В.",
                "report.business_trip_employees": "",
            },
        ],
    )
    book = load_workbook(out)
    sheet = book.active
    assert sheet is not None
    assert sheet["A6"].value == "Продажи"
    assert sheet["A7"].value == "Бухгалтерия"
    assert sheet["B9"].value == "37"
    book.close()


def test_report_templates_guide_marker_catalog_in_sync() -> None:
    guide = _GUIDE.read_text(encoding="utf-8")
    assert marker_catalog_markdown() in guide
