"""Unit: Excel-генерация сводки по филиалу (ADR-0005, новые маркеры)."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from reports.excel_template import (
    TemplateValidationError,
    archive_upload,
    generate_excel_report,
    validate_archived,
)

_SAMPLES = Path(__file__).resolve().parents[2] / "templates_samples"
_BRANCH_SUMMARY = _SAMPLES / "branch_summary_report.xlsx"


def test_branch_summary_template_validates() -> None:
    assert _BRANCH_SUMMARY.is_file()
    validate_archived(_BRANCH_SUMMARY)


def test_branch_summary_unknown_marker_still_rejected(tmp_path: Path) -> None:
    src = tmp_path / "bad.xlsx"
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet["A1"] = "{{отдел_всего}}"
    sheet["A2"] = "{{неизвестный_маркер}}"
    book.save(src)
    book.close()
    with pytest.raises(TemplateValidationError) as exc:
        validate_archived(src)
    assert "неизвестный_маркер" in exc.value.unknown_markers


def test_branch_summary_row_block_expands_multiple_departments(tmp_path: Path) -> None:
    archive = tmp_path / "archive.xlsx"
    archived = archive_upload(_BRANCH_SUMMARY, archive)
    out = tmp_path / "out.xlsx"
    row_records = [
        {
            "employee.division": "Продажи",
            "report.department_total": "25",
            "report.department_absent": "5",
            "report.vacation_employees": "Иванов И.И.; Петров П.П.",
            "report.sick_leave_employees": "Сидоров С.С.",
            "report.business_trip_employees": "Смирнов А.А.; Кузнецов К.К.",
        },
        {
            "employee.division": "Бухгалтерия",
            "report.department_total": "12",
            "report.department_absent": "2",
            "report.vacation_employees": "Орлова О.О.",
            "report.sick_leave_employees": "Волкова В.В.",
            "report.business_trip_employees": "",
        },
    ]
    generate_excel_report(
        archived,
        out,
        scalars={
            "report.title": "Сводка по филиалу",
            "report.date": "19.09.2026",
            "employee.branch": "Центральный",
            "report.branch_total": "37",
            "report.branch_absent": "7",
        },
        row_records=row_records,
    )
    book = load_workbook(out)
    sheet = book.active
    assert sheet is not None
    assert sheet["A1"].value == "Сводка по филиалу"
    assert sheet["A2"].value == "Филиал: Центральный"
    assert sheet["A3"].value == "Дата: 19.09.2026"
    assert sheet["A6"].value == "Продажи"
    assert sheet["B6"].value == "25"
    assert sheet["C6"].value == "5"
    assert sheet["D6"].value == "Иванов И.И.; Петров П.П."
    assert sheet["E6"].value == "Сидоров С.С."
    assert sheet["F6"].value == "Смирнов А.А.; Кузнецов К.К."
    assert sheet["A7"].value == "Бухгалтерия"
    assert sheet["B7"].value == "12"
    assert sheet["C7"].value == "2"
    assert sheet["D7"].value == "Орлова О.О."
    assert sheet["E7"].value == "Волкова В.В."
    assert sheet["F7"].value in ("", None)
    assert sheet["B9"].value == "37"
    assert sheet["C9"].value == "7"
    book.close()
