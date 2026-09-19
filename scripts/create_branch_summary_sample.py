"""Create templates_samples/branch_summary_report.xlsx (one-off generator)."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "templates_samples" / "branch_summary_report.xlsx"


def main() -> None:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Сводка"

    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="D9E1F2")
    title_font = Font(bold=True, size=14)
    header_font = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")
    center = Alignment(horizontal="center", vertical="center")
    right = Alignment(horizontal="right", vertical="center")

    sheet.merge_cells("A1:F1")
    sheet["A1"] = "{{заголовок}}"
    sheet["A1"].font = title_font
    sheet["A1"].alignment = Alignment(horizontal="center")

    sheet["A2"] = "Филиал: {{филиал}}"
    sheet["A3"] = "Дата: {{дата}}"

    headers = (
        "Отдел",
        "Всего",
        "Отсутствуют",
        "Сотрудники в отпуске",
        "Сотрудники на больничном",
        "Сотрудники в командировке",
    )
    markers = (
        "{{отдел}}",
        "{{отдел_всего}}",
        "{{отдел_отсутствуют}}",
        "{{отпуск_ФИО}}",
        "{{больничный_ФИО}}",
        "{{командировка_ФИО}}",
    )
    row_num = 5
    for col, title in enumerate(headers, start=1):
        cell = sheet.cell(row_num, col, title)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = center

    row_num = 6
    for col, marker in enumerate(markers, start=1):
        cell = sheet.cell(row_num, col, marker)
        cell.border = border
        if col in (2, 3):
            cell.alignment = right
        else:
            cell.alignment = wrap
    sheet.cell(row_num, 7, "{{#ROW}}")
    sheet.column_dimensions["G"].hidden = True

    row_num = 8
    sheet["A8"] = "Итого по филиалу:"
    sheet["A8"].font = header_font
    sheet["B8"] = "{{филиал_всего}}"
    sheet["B8"].alignment = right
    sheet["C8"] = "{{филиал_отсутствуют}}"
    sheet["C8"].alignment = right

    widths = {"A": 22, "B": 10, "C": 14, "D": 34, "E": 34, "F": 34}
    for col_letter, width in widths.items():
        sheet.column_dimensions[col_letter].width = width

    sheet.row_dimensions[6].height = 30
    OUT.parent.mkdir(parents=True, exist_ok=True)
    book.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
